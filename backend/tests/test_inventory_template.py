from openpyxl import load_workbook

from app.inventory import COLUMNS, build_template

INVENTORY_SHEET = "IT-Bestand"


def test_template_has_one_header_per_column(tmp_path):
    workbook = load_workbook(build_template(tmp_path / "vorlage.xlsx"))

    headers = [cell.value for cell in workbook[INVENTORY_SHEET][1]]

    assert headers == [column.header for column in COLUMNS]


def test_required_columns_are_marked_in_the_header():
    required = {column.header for column in COLUMNS if column.required}

    assert required == {
        "Dienst *",
        "Zweck *",
        "Kritikalität *",
        "Kritische oder wichtige Funktion *",
    }


def test_every_closed_list_has_a_dropdown(tmp_path):
    workbook = load_workbook(build_template(tmp_path / "vorlage.xlsx"))
    sheet = workbook[INVENTORY_SHEET]
    values = workbook["Werte"]

    validated_columns = {}
    for validation in sheet.data_validations.dataValidation:
        column_letter = str(validation.sqref).split(":")[0].rstrip("0123456789")
        validated_columns[column_letter] = validation.formula1

    for index, column in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        if column.allowed is None:
            assert letter not in validated_columns
            continue
        formula = validated_columns[letter]
        sheet_name, cell_range = formula.split("!")
        first, last = cell_range.replace("$", "").split(":")
        listed = [cell.value for row in values[f"{first}:{last}"] for cell in row]
        assert sheet_name == "Werte"
        assert listed == list(column.allowed)


def test_help_sheet_explains_every_column(tmp_path):
    workbook = load_workbook(build_template(tmp_path / "vorlage.xlsx"))
    help_rows = list(workbook["Ausfüllhilfe"].iter_rows(min_row=2, values_only=True))

    explained = {row[0] for row in help_rows if row[0]}

    assert {column.header for column in COLUMNS} <= explained


def test_dora_columns_name_their_register_field():
    dora = {column.key: column.dora_field for column in COLUMNS if column.dora_field}

    assert dora["service_type"] == "B_02.02.0060"
    assert dora["substitutability"] == "B_07.01.0050"
    assert dora["exit_plan"] == "B_07.01.0080"
    assert all(field.startswith(("B_02.02", "B_06.01", "B_07.01")) for field in dora.values())


def test_service_type_list_follows_the_dora_taxonomy():
    service_type = next(column for column in COLUMNS if column.key == "service_type")

    assert len(service_type.allowed) == 19
    assert [value[:3] for value in service_type.allowed] == [f"S{n:02d}" for n in range(1, 20)]


def test_inventory_sheet_contains_no_example_data(tmp_path):
    """Beispiele stehen in der Ausfüllhilfe, damit sie nie versehentlich importiert werden."""
    workbook = load_workbook(build_template(tmp_path / "vorlage.xlsx"))

    assert workbook[INVENTORY_SHEET].max_row == 1
