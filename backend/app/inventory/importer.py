"""Liest einen IT-Bestand aus einer Excel- oder CSV-Datei.

Die Spalten kommen aus der Definition in template.py. Fehler werden gesammelt und mit
Zeile und Spalte gemeldet, damit der Kunde alle auf einmal korrigieren kann.
"""

import csv
import io
import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from zipfile import BadZipFile

from openpyxl import load_workbook

from app.inventory.template import COLUMNS, INVENTORY_SHEET, Column
from app.models.enums import Criticality

MAX_ROWS = 5000

COLUMNS_BY_KEY = {column.key: column for column in COLUMNS}
INTEGER_COLUMNS = {"notice_period_days"}
DATE_COLUMNS = {"contract_start"}


class UnreadableFile(Exception):
    """Die Datei lässt sich nicht als IT-Bestand lesen."""


@dataclass(frozen=True)
class RowProblem:
    row: int
    column: str | None
    message: str


@dataclass(frozen=True)
class InventoryEntry:
    row: int  # Zeile in der Datei, die Kopfzeile ist Zeile 1
    values: dict[str, object]

    @property
    def service(self) -> str:
        return str(self.values["service"])

    @property
    def criticality(self) -> Criticality:
        return self.values["criticality"]

    @property
    def critical_function(self) -> bool:
        return self.values["critical_function"]


@dataclass
class ImportResult:
    entries: list[InventoryEntry] = field(default_factory=list)
    errors: list[RowProblem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def read_inventory(content: bytes, filename: str) -> ImportResult:
    rows = _read_rows(content, filename)
    if not rows or not any(_is_filled(cell) for cell in rows[0]):
        raise UnreadableFile("Die Datei ist leer oder hat keine Kopfzeile.")
    if len(rows) - 1 > MAX_ROWS:
        raise UnreadableFile(f"Die Datei hat mehr als {MAX_ROWS} Zeilen.")

    result = ImportResult()
    positions = _map_headers(rows[0], result)
    missing = [c for c in COLUMNS if c.required and c.key not in positions]
    for column in missing:
        result.errors.append(RowProblem(1, column.header, "Pflichtspalte fehlt."))
    if missing:
        return result

    for row_number, row in enumerate(rows[1:], start=2):
        if not any(_is_filled(cell) for cell in row):
            continue
        values, errors = _convert_row(row, positions, row_number)
        result.errors.extend(errors)
        if not errors:
            result.entries.append(InventoryEntry(row_number, values))
    return result


def _read_rows(content: bytes, filename: str) -> list[Sequence[object]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, KeyError, OSError, ValueError) as error:
            raise UnreadableFile("Die Excel-Datei lässt sich nicht öffnen.") from error
        sheet = (
            workbook[INVENTORY_SHEET] if INVENTORY_SHEET in workbook.sheetnames else workbook.active
        )
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    if name.endswith(".csv"):
        return _read_csv(content)
    raise UnreadableFile("Erlaubt sind Dateien im Format .xlsx oder .csv.")


def _read_csv(content: bytes) -> list[Sequence[object]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")  # Excel unter Windows speichert CSV oft so
    if not text.strip():
        return []
    first_line = text.splitlines()[0]
    delimiter = ";" if first_line.count(";") >= first_line.count(",") else ","
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def _normalize(header: object) -> str:
    return re.sub(r"\s+", " ", str(header).replace("*", "")).strip().casefold()


def _map_headers(header_row: Iterable[object], result: ImportResult) -> dict[str, int]:
    known = {}
    for column in COLUMNS:
        known[_normalize(column.header)] = column.key
        known[column.key.casefold()] = column.key
    positions = {}
    for index, header in enumerate(header_row):
        if not _is_filled(header):
            continue
        key = known.get(_normalize(header))
        if key is None:
            result.warnings.append(
                f"Spalte „{str(header).strip()}“ ist unbekannt und wird ignoriert."
            )
        elif key in positions:
            result.warnings.append(
                f"Spalte „{str(header).strip()}“ kommt doppelt vor, verwendet wird die erste."
            )
        else:
            positions[key] = index
    return positions


def _convert_row(
    row: Sequence[object], positions: dict[str, int], row_number: int
) -> tuple[dict[str, object], list[RowProblem]]:
    values: dict[str, object] = {}
    errors: list[RowProblem] = []
    for key, index in positions.items():
        column = COLUMNS_BY_KEY[key]
        raw = row[index] if index < len(row) else None
        if not _is_filled(raw):
            if column.required:
                errors.append(RowProblem(row_number, column.header, "Pflichtfeld ist leer."))
            continue
        try:
            values[key] = _convert(column, raw)
        except ValueError as error:
            errors.append(RowProblem(row_number, column.header, str(error)))
    return values, errors


def _convert(column: Column, raw: object) -> object:
    if column.key in INTEGER_COLUMNS:
        return _to_integer(raw)
    if column.key in DATE_COLUMNS:
        return _to_date(raw)
    text = str(raw).strip()
    if column.allowed is None:
        return text
    match = next((value for value in column.allowed if value.casefold() == text.casefold()), None)
    if match is None:
        raise ValueError(f"„{text}“ ist nicht erlaubt. Erlaubt: {', '.join(column.allowed)}.")
    if column.key == "criticality":
        return Criticality[match.upper()]
    if column.key in ("critical_function", "exit_plan"):
        return match == "ja"
    return match


def _to_integer(raw: object) -> int:
    if isinstance(raw, int | float) and float(raw).is_integer() and raw >= 0:
        return int(raw)
    text = str(raw).strip()
    if text.isdigit():
        return int(text)
    raise ValueError(f"„{text}“ ist keine ganze Zahl ab 0.")


def _to_date(raw: object) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    for pattern in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"„{text}“ ist kein Datum im Format TT.MM.JJJJ.")


def _is_filled(value: object) -> bool:
    return value is not None and str(value).strip() != ""
