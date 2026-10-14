"""Vorlage für den IT-Bestand eines Unternehmens.

Die Spalten sind hier einmal definiert. Daraus entstehen die Excel-Vorlage und später der
Import (Schritt 2.3). Wo eine Spalte einem Feld des DORA-Informationsregisters entspricht,
ist das Feld genannt und die Werteliste übernommen (ESA-Dokument JC 2024 79, Annex der
Durchführungsverordnung (EU) 2024/2956). So kann die Datei später auch den Entwurf des
Registers speisen (Phase 10).
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

YES_NO = ("ja", "nein")

# Arten von IKT-Dienstleistungen nach Annex III der Durchführungsverordnung, eigene Übersetzung
SERVICE_TYPES = (
    "S01 IKT-Projektmanagement",
    "S02 IKT-Entwicklung",
    "S03 Helpdesk und First-Level-Support",
    "S04 IKT-Sicherheitsdienste",
    "S05 Bereitstellung von Daten",
    "S06 Datenanalyse",
    "S07 IKT-, Facility- und Hosting-Dienste (ohne Cloud)",
    "S08 Rechenleistung (ohne Cloud)",
    "S09 Datenspeicherung (ohne Cloud)",
    "S10 Telekommunikationsanbieter",
    "S11 Netzwerkinfrastruktur",
    "S12 Hardware und Geräte als Dienst",
    "S13 Softwarelizenzen (ohne SaaS)",
    "S14 IKT-Betrieb einschließlich Wartung",
    "S15 IKT-Beratung",
    "S16 IKT-Risikomanagement",
    "S17 Cloud: IaaS",
    "S18 Cloud: PaaS",
    "S19 Cloud: SaaS",
)


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    required: bool
    help: str
    example: str
    allowed: tuple[str, ...] | None = None
    dora_field: str | None = None
    width: int = 18


COLUMNS = (
    Column(
        "id", "Lfd. Nr.", False, "Eigene Nummer, frei wählbar. Hilft bei Rückfragen.", "17", width=9
    ),
    Column(
        "service",
        "Dienst *",
        True,
        "Name des Dienstes so, wie er im Haus genannt wird. Abkürzungen sind erlaubt, "
        "das Tool ordnet sie zu (z. B. „O365“ oder „Teams“).",
        "Microsoft 365",
        width=28,
    ),
    Column(
        "variant",
        "Variante",
        False,
        "Edition oder Bereitstellungsvariante, falls bekannt (z. B. „EU-Mandant“, „Delos Cloud“). "
        "Ohne Angabe nimmt das Tool die übliche Variante an.",
        "EU-Mandant",
    ),
    Column(
        "provider",
        "Vertragspartner",
        False,
        "Unternehmen, mit dem der Vertrag besteht. Erleichtert die Zuordnung.",
        "Microsoft Ireland Operations Ltd.",
        width=28,
    ),
    Column(
        "purpose",
        "Zweck *",
        True,
        "Welche Geschäftsfunktion der Dienst unterstützt.",
        "E-Mail, Dokumente und Videokonferenzen für alle Mitarbeitenden",
        dora_field="B_06.01.0030",
        width=36,
    ),
    Column("department", "Abteilung", False, "Fachlich verantwortliche Abteilung.", "IT"),
    Column(
        "criticality",
        "Kritikalität *",
        True,
        "Wie schwer wöge ein Ausfall für das Unternehmen? Bestimmt das Gewicht des Dienstes "
        "in der Gesamtbewertung (gering 1, mittel 2, hoch 3).",
        "hoch",
        allowed=("gering", "mittel", "hoch"),
        width=14,
    ),
    Column(
        "critical_function",
        "Kritische oder wichtige Funktion *",
        True,
        "Unterstützt der Dienst eine kritische oder wichtige Funktion im Sinne von DORA? "
        "Für Unternehmen außerhalb des Finanzsektors „nein“ wählen.",
        "ja",
        allowed=YES_NO,
        width=16,
    ),
    Column(
        "data_types",
        "Datenarten",
        False,
        "Welche Daten der Dienst verarbeitet, mit Semikolon getrennt "
        "(z. B. Kundendaten; Personaldaten; Zahlungsdaten).",
        "Kundendaten; Personaldaten",
        width=26,
    ),
    Column(
        "data_sensitivity",
        "Sensibilität der Daten",
        False,
        "Die sensibelsten Daten geben den Ausschlag.",
        "hoch",
        allowed=("niedrig", "mittel", "hoch"),
        dora_field="B_02.02.0170",
        width=14,
    ),
    Column(
        "service_type",
        "Art der IKT-Dienstleistung",
        False,
        "Einordnung nach der Liste S01 bis S19 aus Annex III der Durchführungsverordnung.",
        "S19 Cloud: SaaS",
        allowed=SERVICE_TYPES,
        dora_field="B_02.02.0060",
        width=30,
    ),
    Column(
        "contract_reference",
        "Vertragsreferenz",
        False,
        "Vertrags- oder Bestellnummer.",
        "V-2023-0042",
        dora_field="B_02.02.0010",
    ),
    Column(
        "contract_start",
        "Vertragsbeginn",
        False,
        "Datum im Format TT.MM.JJJJ.",
        "01.04.2023",
        dora_field="B_02.02.0070",
        width=14,
    ),
    Column(
        "notice_period_days",
        "Kündigungsfrist (Tage)",
        False,
        "Kündigungsfrist für das eigene Unternehmen in Kalendertagen.",
        "90",
        dora_field="B_02.02.0100",
        width=14,
    ),
    Column(
        "reliance",
        "Abhängigkeitsgrad",
        False,
        "Wie stark hängt die Funktion an diesem Dienst? „wesentlich“: spürbare Schäden nach "
        "Minuten bis Stunden. „vollständig“: sofortige, schwere und lange Unterbrechung.",
        "vollständig",
        allowed=("nicht signifikant", "gering", "wesentlich", "vollständig"),
        dora_field="B_02.02.0180",
        width=16,
    ),
    Column(
        "substitutability",
        "Ersetzbarkeit",
        False,
        "Wie schwer ließe sich der Anbieter ersetzen?",
        "sehr komplex",
        allowed=("nicht ersetzbar", "sehr komplex", "mittel", "leicht"),
        dora_field="B_07.01.0050",
        width=16,
    ),
    Column(
        "exit_plan",
        "Exit-Plan vorhanden",
        False,
        "Gibt es einen dokumentierten Plan für den Ausstieg aus diesem Dienst?",
        "nein",
        allowed=YES_NO,
        dora_field="B_07.01.0080",
        width=12,
    ),
    Column(
        "reintegration",
        "Rückholung möglich",
        False,
        "Ließe sich der Dienst ins eigene Haus zurückholen?",
        "schwierig",
        allowed=("leicht", "schwierig", "sehr komplex"),
        dora_field="B_07.01.0090",
        width=14,
    ),
    Column("notes", "Bemerkungen", False, "Alles, was sonst wichtig ist.", "", width=30),
)

INVENTORY_SHEET = "IT-Bestand"
HELP_SHEET = "Ausfüllhilfe"
VALUES_SHEET = "Werte"
INPUT_ROWS = 1000  # so viele Zeilen erhalten Auswahllisten

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_template(path: Path) -> Path:
    workbook = Workbook()
    inventory = workbook.active
    inventory.title = INVENTORY_SHEET
    help_sheet = workbook.create_sheet(HELP_SHEET)
    values = workbook.create_sheet(VALUES_SHEET)

    _fill_inventory(inventory, values)
    _fill_help(help_sheet)
    values.sheet_state = "hidden"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _fill_inventory(sheet, values) -> None:
    value_column = 0
    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        cell = sheet.cell(row=1, column=index, value=column.header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[letter].width = column.width

        if column.allowed is None:
            continue
        value_column += 1
        value_letter = get_column_letter(value_column)
        values.cell(row=1, column=value_column, value=column.header)
        for row, option in enumerate(column.allowed, start=2):
            values.cell(row=row, column=value_column, value=option)
        last = len(column.allowed) + 1
        validation = DataValidation(
            type="list",
            formula1=f"{VALUES_SHEET}!${value_letter}$2:${value_letter}${last}",
            allow_blank=not column.required,
            showErrorMessage=True,
            errorTitle="Ungültiger Wert",
            error="Bitte einen Wert aus der Liste wählen.",
        )
        validation.add(f"{letter}2:{letter}{INPUT_ROWS + 1}")
        sheet.add_data_validation(validation)

    sheet.row_dimensions[1].height = 45
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def _fill_help(sheet) -> None:
    headers = ("Spalte", "Pflicht", "Erklärung", "Erlaubte Werte", "Beispiel", "DORA-Register")
    widths = (32, 9, 70, 40, 34, 16)
    for index, (header, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row, column in enumerate(COLUMNS, start=2):
        entries = (
            column.header,
            "ja" if column.required else "nein",
            column.help,
            "; ".join(column.allowed) if column.allowed else "",
            column.example,
            column.dora_field or "",
        )
        for index, value in enumerate(entries, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    note_row = len(COLUMNS) + 3
    notes = (
        "So geht es: Im Blatt „IT-Bestand“ je Dienst eine Zeile ausfüllen. "
        "Spalten mit * sind Pflicht, alle anderen helfen bei einer genaueren Bewertung.",
        "Datenstandort, Konzernmutter, Vertragsrecht und Zertifikate müssen nicht eingetragen "
        "werden. Sie stammen aus der Anbieterdatenbank des Tools.",
        "Die Spalten mit DORA-Feld entsprechen dem Informationsregister nach der "
        "Durchführungsverordnung (EU) 2024/2956. Die Wertelisten sind daraus übernommen und "
        "übersetzt.",
    )
    for offset, note in enumerate(notes):
        cell = sheet.cell(row=note_row + offset, column=1, value=note)
        sheet.merge_cells(
            start_row=note_row + offset, start_column=1, end_row=note_row + offset, end_column=6
        )
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[note_row + offset].height = 32
    sheet.freeze_panes = "A2"
